from pathlib import Path

import pytest
from docx import Document
import fitz

from app.schemas.resume import ResumeMetadata
from app.services.resume_parser.docx_parser import DOCXParser
from app.services.resume_parser.pdf_parser import PDFParser


def test_pdf_parser_extracts_metadata(tmp_path: Path) -> None:
    pdf_path = tmp_path / "sample.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Jane Doe\n jane@example.com \n 555-123-4567\n Seattle, WA\n Skills: Python, React")
    document.save(pdf_path)
    document.close()

    text, metadata = PDFParser().parse(pdf_path)

    assert "Jane Doe" in text
    assert metadata.name == "Jane Doe"
    assert metadata.email == "jane@example.com"
    assert metadata.phone == "555-123-4567"
    assert metadata.location == "Seattle, WA"
    assert "Python" in metadata.skills


def test_docx_parser_extracts_metadata(tmp_path: Path) -> None:
    doc_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("jane@example.com")
    document.add_paragraph("Skills: Python, React")
    document.save(doc_path)

    text, metadata = DOCXParser().parse(doc_path)

    assert "Jane Doe" in text
    assert metadata.name == "Jane Doe"
    assert metadata.email == "jane@example.com"
    assert "Python" in metadata.skills


def test_resume_metadata_accepts_optional_fields() -> None:
    metadata = ResumeMetadata(name="Jane Doe", skills=["Python"])

    assert metadata.name == "Jane Doe"
    assert metadata.skills == ["Python"]
